"""نظام إدارة حصص الاحتياط — قسم مستقل بحساب دخول خاص."""
import base64
import io
import json
import os
import random
import re
import uuid
from datetime import datetime, timezone, timedelta, date as ddate
from pathlib import Path
from typing import List, Optional

import jwt
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse, Response
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel

from substitution_parser import parse_teacher_pdf, parse_general_pdf, parse_timing_text, parse_supervision_pdf, DAYS

ROOT = Path(__file__).parent
SEED_FILE = ROOT / "data" / "substitution_seed.json"
SCHOOL_NAME = "مدرسة الخيرات للبنين ٥-٨"
WEEKDAY_TO_DAY = {6: "الأحد", 0: "الاثنين", 1: "الثلاثاء", 2: "الأربعاء", 3: "الخميس"}
DEFAULT_PERIOD_TIMES = [
    {"from": "7:25", "to": "8:05"}, {"from": "8:05", "to": "8:45"}, {"from": "8:45", "to": "9:25"}, {"from": "9:25", "to": "10:05"},
    {"from": "10:35", "to": "11:15"}, {"from": "11:15", "to": "11:55"}, {"from": "11:55", "to": "12:35"}, {"from": "1:00", "to": "1:40"},
]
BUNDLED_TIMING = ROOT / "data" / "school_timing.jpeg"
security = HTTPBearer()


def _key(name: str) -> str:
    return name.replace(" ", "").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ى", "ي").replace("ّ", "")


def _tokens(name: str):
    toks = [t for t in re.split(r"\s+", name.strip()) if t]
    out = []
    for t in toks:
        k = _key(t).replace("ؤ", "و")
        if k in ("بن", "بنت"):
            continue
        if out and out[-1] in ("عبد", "ابو"):
            out[-1] += k
        else:
            out.append(k)
    return out, "".join(out)


def _is_subseq(a, b):
    it = iter(b)
    return all(any(x == y for y in it) for x in a)


def match_teacher(name: str, tlist):
    """يطابق اسم مشرف من جدول الإشراف مع معلم من قائمة المعلمين (مطابقة تامة ثم تسلسل الأسماء ثم احتواء)."""
    st, sk = _tokens(name)
    if not st:
        return None
    best, best_rank = None, None
    for t in tlist:
        tt, tk = _tokens(t["name"])
        if not tt or st[0] != tt[0]:
            continue
        if sk == tk:
            rank = (0, 0)
        elif _is_subseq(st, tt) or _is_subseq(tt, st):
            rank = (1, abs(len(tt) - len(st)))
        elif len(min(sk, tk, key=len)) >= 6 and (sk in tk or tk in sk):
            rank = (2, abs(len(tk) - len(sk)))
        else:
            continue
        if best_rank is None or rank < best_rank:
            best, best_rank = t, rank
    return best


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def day_name_of(d: str) -> str:
    return WEEKDAY_TO_DAY.get(ddate.fromisoformat(d).weekday(), "")


def academic_range(d: str):
    dt = ddate.fromisoformat(d)
    y = dt.year if dt.month >= 9 else dt.year - 1
    return f"{y}-09-01", f"{y + 1}-08-31"


def _dur(p) -> int:
    try:
        h1, m1 = map(int, p["from"].split(":")); h2, m2 = map(int, p["to"].split(":"))
        if h1 < 6: h1 += 12
        if h2 < 6: h2 += 12
        return (h2 * 60 + m2) - (h1 * 60 + m1)
    except Exception:
        return -1


def fmt_date_ar(d: str) -> str:
    dt = ddate.fromisoformat(d)
    return f"{dt.day}/{dt.month}/{dt.year}"


class LoginReq(BaseModel):
    username: str
    password: str


class ChangePwdReq(BaseModel):
    old_password: str
    new_password: str


class AbsentReq(BaseModel):
    teacher_ids: List[str]


class AssignReq(BaseModel):
    absent_id: str
    period: int
    substitute_id: Optional[str] = None


class TeacherUpdate(BaseModel):
    name: Optional[str] = None
    subject: Optional[str] = None
    quota: Optional[int] = None
    active: Optional[bool] = None


class TeacherCreate(BaseModel):
    name: str
    subject: str = ""
    quota: int = 0


class AccountReq(BaseModel):
    current_password: str
    username: Optional[str] = None
    new_password: Optional[str] = None


class SettingsReq(BaseModel):
    school_name: Optional[str] = None
    period_times: Optional[List[dict]] = None


class SupervisionDay(BaseModel):
    day: str
    leader: str = ""
    supervisors: List[str] = []


class SupervisionReq(BaseModel):
    days: List[SupervisionDay]


class BulkIdsReq(BaseModel):
    ids: List[str]


class BulkActiveReq(BaseModel):
    ids: List[str]
    active: bool


class SwapCreateReq(BaseModel):
    teacher_a_id: str
    period_a: int
    teacher_b_id: str
    period_b: int


def make_router(db, hash_password, verify_password, make_token, jwt_secret, jwt_algorithm):
    router = APIRouter(prefix="/substitution")
    users = db.sub_users
    teachers = db.sub_teachers
    days = db.sub_days
    settings = db.sub_settings
    supervision = db.sub_supervision
    swaps = db.sub_swaps

    async def load_supervision():
        doc = await supervision.find_one({"id": "main"}, {"_id": 0})
        if not doc:
            doc = {"id": "main", "days": [{"day": d, "leader": "", "supervisors": []} for d in DAYS], "updated_at": None}
        return doc

    def resolve_supervision(doc, tlist):
        """يُرجع أيام الإشراف مع ربط كل اسم بمعلم من القائمة إن وُجد."""
        out = []
        for d in doc["days"]:
            leader = match_teacher(d["leader"], tlist) if d.get("leader") else None
            sups = [{"name": n, "teacher_id": (m or {}).get("id")} for n in d.get("supervisors", []) for m in [match_teacher(n, tlist)]]
            out.append({"day": d["day"], "leader": {"name": d.get("leader", ""), "teacher_id": (leader or {}).get("id")}, "supervisors": sups})
        return out

    def supervisors_of_day(resolved, day_name):
        """{teacher_id: 'leader'|'member'} لليوم المحدد."""
        for d in resolved:
            if d["day"] == day_name:
                roles = {s["teacher_id"]: "member" for s in d["supervisors"] if s["teacher_id"]}
                if d["leader"]["teacher_id"]:
                    roles[d["leader"]["teacher_id"]] = "leader"
                return roles
        return {}

    async def get_settings():
        doc = await settings.find_one({"id": "main"}, {"_id": 0, "timing_data": 0})
        if not doc:
            doc = {"id": "main", "school_name": SCHOOL_NAME, "period_times": DEFAULT_PERIOD_TIMES, "timing_file": None, "updated_at": now_iso()}
            if BUNDLED_TIMING.exists():
                doc.update({"timing_file": "/api/substitution/settings/timing-file?v=seed", "timing_mime": "image/jpeg",
                            "timing_data": base64.b64encode(BUNDLED_TIMING.read_bytes()).decode()})
            await settings.insert_one(dict(doc))
            doc.pop("timing_data", None)
        pts = doc.get("period_times") or DEFAULT_PERIOD_TIMES
        doc["period_times"] = (pts + DEFAULT_PERIOD_TIMES)[:8]
        doc["period_labels"] = [f"{p.get('from', '')} - {p.get('to', '')}" for p in doc["period_times"]]
        return doc

    async def init():
        await teachers.create_index("name")
        await days.create_index("date", unique=True)
        u = os.environ.get("SUB_ADMIN_USERNAME")
        p = os.environ.get("SUB_ADMIN_PASSWORD")
        if u and p:
            existing = await users.find_one({"username": u})
            if not existing:
                await users.insert_one({"id": str(uuid.uuid4()), "username": u, "password_hash": hash_password(p), "name": "إدارة الاحتياط", "created_at": now_iso()})
        if await teachers.count_documents({}) == 0 and SEED_FILE.exists():
            seed = json.loads(SEED_FILE.read_text(encoding="utf-8"))
            docs = [{"id": str(uuid.uuid4()), "order": i, "active": True, **t} for i, t in enumerate(seed)]
            if docs:
                await teachers.insert_many(docs)

    router.init = init

    async def current_user(creds: HTTPAuthorizationCredentials = Depends(security)):
        try:
            payload = jwt.decode(creds.credentials, jwt_secret, algorithms=[jwt_algorithm])
        except jwt.ExpiredSignatureError:
            raise HTTPException(status_code=401, detail="انتهت صلاحية الجلسة")
        except jwt.InvalidTokenError:
            raise HTTPException(status_code=401, detail="جلسة غير صالحة")
        if payload.get("role") != "substitution":
            raise HTTPException(status_code=401, detail="هذه الجلسة ليست لنظام الاحتياط")
        user = await users.find_one({"id": payload.get("sub_user_id")}, {"_id": 0, "password_hash": 0})
        if not user:
            raise HTTPException(status_code=401, detail="الحساب غير موجود")
        return user

    # ---------------- auth ----------------
    @router.post("/auth/login")
    async def login(req: LoginReq):
        username = req.username.strip().lower()
        key = f"sub:{username}"
        att = await db.login_attempts.find_one({"identifier": key})
        if att and att.get("locked_until") and datetime.fromisoformat(att["locked_until"]) > datetime.now(timezone.utc):
            raise HTTPException(status_code=429, detail="تم قفل تسجيل الدخول مؤقتاً بسبب محاولات خاطئة متكررة، حاول بعد 15 دقيقة")
        doc = await users.find_one({"username": username})
        if not doc or not verify_password(req.password, doc["password_hash"]):
            count = (att or {}).get("count", 0) + 1
            upd = {"count": count, "locked_until": None}
            if count >= 5:
                upd = {"count": 0, "locked_until": (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()}
            await db.login_attempts.update_one({"identifier": key}, {"$set": upd}, upsert=True)
            raise HTTPException(status_code=401, detail="اسم المستخدم أو كلمة المرور غير صحيحة")
        await db.login_attempts.delete_one({"identifier": key})
        return {"token": make_token({"role": "substitution", "sub_user_id": doc["id"], "username": username}), "name": doc.get("name", ""), "username": username}

    @router.get("/auth/me")
    async def me(u=Depends(current_user)):
        return u

    @router.post("/auth/change-password")
    async def change_password(req: ChangePwdReq, u=Depends(current_user)):
        doc = await users.find_one({"id": u["id"]})
        if not verify_password(req.old_password, doc["password_hash"]):
            raise HTTPException(status_code=400, detail="كلمة المرور القديمة غير صحيحة")
        if len(req.new_password) < 6:
            raise HTTPException(status_code=400, detail="كلمة المرور الجديدة قصيرة (6 أحرف على الأقل)")
        await users.update_one({"id": u["id"]}, {"$set": {"password_hash": hash_password(req.new_password)}})
        return {"message": "تم تغيير كلمة المرور"}

    @router.put("/auth/account")
    async def update_account(req: AccountReq, u=Depends(current_user)):
        doc = await users.find_one({"id": u["id"]})
        if not verify_password(req.current_password, doc["password_hash"]):
            raise HTTPException(status_code=400, detail="كلمة المرور الحالية غير صحيحة")
        upd = {}
        if req.username and req.username.strip().lower() != doc["username"]:
            new_u = req.username.strip().lower()
            if len(new_u) < 3:
                raise HTTPException(status_code=400, detail="اسم المستخدم قصير (3 أحرف على الأقل)")
            if await users.find_one({"username": new_u}):
                raise HTTPException(status_code=400, detail="اسم المستخدم مستخدم بالفعل")
            upd["username"] = new_u
        if req.new_password:
            if len(req.new_password) < 6:
                raise HTTPException(status_code=400, detail="كلمة المرور الجديدة قصيرة (6 أحرف على الأقل)")
            upd["password_hash"] = hash_password(req.new_password)
        if not upd:
            raise HTTPException(status_code=400, detail="لا يوجد تغيير")
        await users.update_one({"id": u["id"]}, {"$set": upd})
        username = upd.get("username", doc["username"])
        return {"message": "تم تحديث الحساب", "username": username,
                "token": make_token({"role": "substitution", "sub_user_id": doc["id"], "username": username})}

    # ---------------- settings ----------------
    @router.get("/settings")
    async def read_settings(u=Depends(current_user)):
        return await get_settings()

    @router.put("/settings")
    async def write_settings(req: SettingsReq, u=Depends(current_user)):
        upd = {"updated_at": now_iso()}
        if req.school_name is not None:
            upd["school_name"] = req.school_name.strip() or SCHOOL_NAME
        if req.period_times is not None:
            upd["period_times"] = [{"from": str(p.get("from", "")).strip(), "to": str(p.get("to", "")).strip()} for p in req.period_times][:8]
        await get_settings()
        await settings.update_one({"id": "main"}, {"$set": upd})
        return await get_settings()

    async def save_timing_file(file: UploadFile):
        ext = (file.filename or "").rsplit(".", 1)[-1].lower()
        mimes = {"pdf": "application/pdf", "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}
        if ext not in mimes:
            raise HTTPException(status_code=400, detail="ملف التوقيت يجب أن يكون صورة (JPG/PNG) أو PDF")
        raw = await file.read()
        if len(raw) > 6 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="حجم ملف التوقيت كبير (الحد 6MB)")
        detected = []
        if ext == "pdf":
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(raw)) as pdf:
                    text = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
                detected = parse_timing_text(text)
            except Exception:
                detected = []
        ver = uuid.uuid4().hex[:6]
        await get_settings()
        await settings.update_one({"id": "main"}, {"$set": {
            "timing_file": f"/api/substitution/settings/timing-file?v={ver}", "timing_mime": mimes[ext],
            "timing_data": base64.b64encode(raw).decode(), "updated_at": now_iso()}})
        return detected

    @router.get("/settings/timing-file")
    async def timing_file():
        doc = await settings.find_one({"id": "main"}, {"_id": 0, "timing_data": 1, "timing_mime": 1})
        if not doc or not doc.get("timing_data"):
            raise HTTPException(status_code=404, detail="لا يوجد ملف توقيت")
        return Response(content=base64.b64decode(doc["timing_data"]), media_type=doc.get("timing_mime", "image/jpeg"))

    @router.post("/settings/timing")
    async def upload_timing(file: UploadFile = File(...), u=Depends(current_user)):
        detected = await save_timing_file(file)
        out = await get_settings()
        out["detected_times"] = detected
        return out

    # ---------------- helpers ----------------
    async def load_counts(from_d: str, to_d: str):
        """عدد أيام الغياب وحصص الاحتياط لكل معلم ضمن فترة."""
        abs_c, sub_c = {}, {}
        async for d in days.find({"date": {"$gte": from_d, "$lte": to_d}}, {"_id": 0}):
            for tid in d.get("absent", []):
                abs_c[tid] = abs_c.get(tid, 0) + 1
            for a in d.get("assignments", []):
                sid = a.get("substitute_id")
                if sid:
                    sub_c[sid] = sub_c.get(sid, 0) + 1
        return abs_c, sub_c

    async def load_period_subs(from_d: str, to_d: str, period: int):
        """عدد حصص الاحتياط في حصة معينة (مثل الحصة الثامنة) لكل معلم ضمن فترة."""
        c = {}
        async for d in days.find({"date": {"$gte": from_d, "$lte": to_d}}, {"_id": 0, "assignments": 1}):
            for a in d.get("assignments", []):
                sid = a.get("substitute_id")
                if sid and a.get("period") == period:
                    c[sid] = c.get(sid, 0) + 1
        return c

    async def get_or_new_day(d: str):
        doc = await days.find_one({"date": d}, {"_id": 0})
        if not doc:
            doc = {"date": d, "day_name": day_name_of(d), "absent": [], "assignments": [], "late_ids": [], "updated_at": now_iso()}
        doc.setdefault("late_ids", [])
        return doc

    async def save_day(doc):
        doc["updated_at"] = now_iso()
        await days.update_one({"date": doc["date"]}, {"$set": doc}, upsert=True)

    def slot_of(t, day_name, period):
        sched = t.get("schedule", {}).get(day_name) or [None] * 8
        return sched[period - 1] if 1 <= period <= 8 else None

    def prev_school_date(d: str) -> str:
        cur = ddate.fromisoformat(d)
        while True:
            cur -= timedelta(days=1)
            if WEEKDAY_TO_DAY.get(cur.weekday()):
                return cur.isoformat()

    async def prev_sub_sets(d: str, n: int = 3):
        """مجموعات معرّفات المعلمين الذين كُلِّفوا باحتياط في آخر n أيام دراسية سابقة (0=الأمس)."""
        out = []
        cur = d
        for _ in range(n):
            cur = prev_school_date(cur)
            doc = await days.find_one({"date": cur}, {"_id": 0, "assignments": 1})
            out.append({a.get("substitute_id") for a in (doc or {}).get("assignments", []) if a.get("substitute_id")})
        return out

    def teaches_class(t, klass):
        if not klass:
            return False
        for day_cells in t.get("schedule", {}).values():
            for c in day_cells:
                if c and c.get("class") == klass:
                    return True
        return False

    async def rank_candidates(doc, tmap, absent_id, period, exclude_busy=True, exclude_supervisors=False):
        d = doc["date"]
        day_name = doc["day_name"]
        y0, y1 = academic_range(d)
        _, sub_year = await load_counts(y0, y1)
        p8_subs = await load_period_subs(y0, y1, 8)
        sup_roles = supervisors_of_day(resolve_supervision(await load_supervision(), list(tmap.values())), day_name)
        today_subs = {}
        taken_this_period = set()
        for a in doc["assignments"]:
            if a.get("substitute_id"):
                today_subs[a["substitute_id"]] = today_subs.get(a["substitute_id"], 0) + 1
                if a["period"] == period:
                    taken_this_period.add(a["substitute_id"])
        absent_t = tmap.get(absent_id) or {}
        target_class = (slot_of(absent_t, day_name, period) or {}).get("class")
        active_quotas = sorted((t.get("quota", 0) for t in tmap.values() if t.get("active", True)), reverse=True)
        hq_threshold = active_quotas[max(0, int(len(active_quotas) * 0.2) - 1)] if active_quotas else 0
        streak_sets = await prev_sub_sets(d, 3)
        out = []
        for t in tmap.values():
            if not t.get("active", True) or t["id"] in doc["absent"]:
                continue
            role = sup_roles.get(t["id"])
            if exclude_supervisors and role:
                continue
            slot = slot_of(t, day_name, period)
            free = slot is None and t["id"] not in taken_this_period
            if exclude_busy and not free:
                continue
            day_periods = sum(1 for c in (t.get("schedule", {}).get(day_name) or []) if c)
            consecutive_alert = t["id"] in streak_sets[0] and t["id"] in streak_sets[1] and t["id"] not in streak_sets[2]
            out.append({
                "id": t["id"], "name": t["name"], "subject": t.get("subject", ""), "quota": t.get("quota", 0),
                "subs_year": sub_year.get(t["id"], 0), "subs_today": today_subs.get(t["id"], 0),
                "subs_year_p8": p8_subs.get(t["id"], 0),
                "free": free, "busy_class": (slot or {}).get("class") if slot else None,
                "same_class": teaches_class(t, target_class),
                "already_taken": t["id"] in taken_this_period,
                "high_quota": hq_threshold > 0 and t.get("quota", 0) >= hq_threshold,
                "day_periods": day_periods,
                "consecutive_alert": consecutive_alert,
                "supervisor": role,
                "schedule": t.get("schedule", {}),
            })
        out.sort(key=lambda c: (not c["free"], c["quota"], not c["same_class"], c["subs_today"], c["subs_year"], c["name"]))
        if out:
            frees = [c for c in out if c["free"]]
            if frees:
                min_q = min(c["quota"] for c in frees)
                min_s = min(c["subs_year"] for c in frees)
                for c in frees:
                    c["least_quota"] = c["quota"] == min_q
                    c["least_subs"] = c["subs_year"] == min_s
        return out

    def pick_fair_candidate(cands):
        """يختار بديلاً بعدالة (الأولوية للأقل نصاباً) مع تنويع النتيجة بين كل تشغيل للتوزيع التلقائي."""
        free = [c for c in cands if c["free"]]
        if not free:
            return None
        top_quota = free[0]["quota"]
        pool = [c for c in free if c["quota"] == top_quota]
        if len(pool) < 3:
            near = [c for c in free if c not in pool and c["quota"] - top_quota <= 3]
            pool += near[: max(0, 3 - len(pool))]
        return random.choice(pool)

    async def enrich_day(doc):
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        cfg = await get_settings()
        PERIOD_TIMES = cfg["period_labels"]
        y0, y1 = academic_range(doc["date"])
        abs_year, sub_year = await load_counts(y0, y1)
        absent = []
        for tid in doc["absent"]:
            t = tmap.get(tid)
            if not t:
                continue
            sched = t.get("schedule", {}).get(doc["day_name"]) or [None] * 8
            periods = []
            for i, cell in enumerate(sched):
                a = next((x for x in doc["assignments"] if x["absent_id"] == tid and x["period"] == i + 1), None)
                periods.append({"period": i + 1, "time": PERIOD_TIMES[i], "class": (cell or {}).get("class"), "subject": (cell or {}).get("subject"),
                                "substitute_id": (a or {}).get("substitute_id"), "substitute_name": tmap.get((a or {}).get("substitute_id"), {}).get("name")})
            absent.append({"id": tid, "name": t["name"], "subject": t.get("subject", ""), "quota": t.get("quota", 0),
                           "absences_year": abs_year.get(tid, 0), "periods": periods, "late": tid in doc.get("late_ids", [])})
        rows = []
        order = {tid: i for i, tid in enumerate(doc["absent"])}
        resolved = resolve_supervision(await load_supervision(), list(tmap.values()))
        sup_roles = supervisors_of_day(resolved, doc["day_name"])
        for a in sorted(doc["assignments"], key=lambda x: (order.get(x["absent_id"], 99), x["period"])):
            at, st = tmap.get(a["absent_id"], {}), tmap.get(a.get("substitute_id"), {})
            rows.append({**a, "subject": a.get("subject") or at.get("subject", ""), "absent_name": at.get("name", ""), "substitute_name": st.get("name", ""), "substitute_subject": st.get("subject", ""),
                         "substitute_supervisor": sup_roles.get(a.get("substitute_id")), "time": PERIOD_TIMES[a["period"] - 1],
                         "substitute_subs_year": sub_year.get(a.get("substitute_id"), 0)})
        is_school = bool(doc["day_name"])
        sup_day = next((x for x in resolved if x["day"] == doc["day_name"]), None)
        return {"date": doc["date"], "day_name": doc["day_name"], "is_school_day": is_school, "date_ar": fmt_date_ar(doc["date"]),
                "absent": absent, "assignments": rows, "school_name": cfg["school_name"], "period_times": PERIOD_TIMES, "updated_at": doc.get("updated_at"),
                "supervision": sup_day}

    # ---------------- teachers ----------------
    @router.get("/teachers")
    async def list_teachers(date: Optional[str] = None, u=Depends(current_user)):
        d = date or ddate.today().isoformat()
        y0, y1 = academic_range(d)
        abs_c, sub_c = await load_counts(y0, y1)
        out = []
        async for t in teachers.find({}, {"_id": 0}).sort("order", 1):
            out.append({**t, "absences": abs_c.get(t["id"], 0), "subs": sub_c.get(t["id"], 0)})
        cfg = await get_settings()
        return {"teachers": out, "academic_range": [y0, y1], "days": DAYS, "period_times": cfg["period_labels"]}

    @router.post("/teachers")
    async def create_teacher(req: TeacherCreate, u=Depends(current_user)):
        n = await teachers.count_documents({})
        doc = {"id": str(uuid.uuid4()), "name": req.name.strip(), "subject": req.subject.strip(), "quota": req.quota,
               "schedule": {d: [None] * 8 for d in DAYS}, "active": True, "order": n}
        await teachers.insert_one(doc)
        doc.pop("_id", None)
        return doc

    @router.put("/teachers/bulk-active")
    async def bulk_active(req: BulkActiveReq, u=Depends(current_user)):
        r = await teachers.update_many({"id": {"$in": req.ids}}, {"$set": {"active": req.active}})
        return {"updated": r.modified_count}

    @router.put("/teachers/{tid}")
    async def update_teacher(tid: str, req: TeacherUpdate, u=Depends(current_user)):
        upd = {k: v for k, v in req.model_dump().items() if v is not None}
        if "name" in upd:
            upd["name"] = upd["name"].strip()
        r = await teachers.update_one({"id": tid}, {"$set": upd})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="المعلم غير موجود")
        return await teachers.find_one({"id": tid}, {"_id": 0})

    @router.delete("/teachers/all")
    async def delete_all_teachers(u=Depends(current_user)):
        n = await teachers.count_documents({})
        await teachers.delete_many({})
        return {"message": "تم حذف جميع المعلمين", "deleted": n}

    @router.post("/teachers/bulk-delete")
    async def bulk_delete(req: BulkIdsReq, u=Depends(current_user)):
        r = await teachers.delete_many({"id": {"$in": req.ids}})
        return {"deleted": r.deleted_count}

    @router.delete("/teachers/{tid}")
    async def delete_teacher(tid: str, u=Depends(current_user)):
        await teachers.delete_one({"id": tid})
        return {"ok": True}

    @router.post("/teachers/import")
    async def import_files(teachers_pdf: Optional[UploadFile] = File(None), general_pdf: Optional[UploadFile] = File(None),
                           timing: Optional[UploadFile] = File(None), file: Optional[UploadFile] = File(None), u=Depends(current_user)):
        teachers_pdf = teachers_pdf or file
        if not teachers_pdf and not general_pdf and not timing:
            raise HTTPException(status_code=400, detail="اختر ملفاً واحداً على الأقل للاستيراد")

        async def read_pdf(up, parser, label):
            if not (up.filename or "").lower().endswith(".pdf"):
                raise HTTPException(status_code=400, detail=f"{label}: يجب أن يكون ملف PDF")
            tmp = ROOT / "data" / f"_import_{uuid.uuid4().hex}.pdf"
            tmp.write_bytes(await up.read())
            try:
                return parser(str(tmp))
            except Exception:
                raise HTTPException(status_code=400, detail=f"{label}: تعذّر قراءة الملف — تأكد أنه صادر من aSc Timetables")
            finally:
                tmp.unlink(missing_ok=True)

        result = {"added": 0, "updated": 0, "deactivated": 0, "total": 0, "quota_mismatch": [], "timing_detected": 0, "sources": []}
        parsed_t = await read_pdf(teachers_pdf, parse_teacher_pdf, "جدول حصص المعلمين") if teachers_pdf else []
        parsed_g = await read_pdf(general_pdf, parse_general_pdf, "الجدول العام") if general_pdf else []
        if teachers_pdf and not parsed_t:
            raise HTTPException(status_code=400, detail="جدول حصص المعلمين: لم يُعثر على أي جدول معلم")
        if general_pdf and not parsed_g:
            raise HTTPException(status_code=400, detail="الجدول العام: لم يُعثر على أسماء المعلمين وأنصبتهم")
        if parsed_t:
            result["sources"].append("جدول حصص المعلمين")
        if parsed_g:
            result["sources"].append("الجدول العام")

        # الدمج: الجدول التفصيلي هو الأساس (يحوي المواد)، والعام يُكمل الناقص ويتحقق من الأنصبة
        merged = {}
        for t in parsed_t:
            merged[_key(t["name"])] = dict(t)
        for g in parsed_g:
            k = _key(g["name"])
            if k in merged:
                if merged[k]["quota"] != g["quota"]:
                    result["quota_mismatch"].append({"name": g["name"], "teachers_pdf": merged[k]["quota"], "general_pdf": g["quota"]})
                    merged[k]["quota"] = g["quota"]
            else:
                merged[k] = {"name": g["name"], "subject": "", "quota": g["quota"], "schedule": g["schedule"]}

        if merged:
            existing = {_key(t["name"]): t async for t in teachers.find({}, {"_id": 0})}
            seen = set()
            for i, (k, t) in enumerate(merged.items()):
                seen.add(k)
                subj = t.get("subject") or (existing.get(k) or {}).get("subject", "")
                for cells in t["schedule"].values():
                    for c in cells:
                        if c and not c.get("subject"):
                            c["subject"] = subj
                if k in existing:
                    upd = {"quota": t["quota"], "schedule": t["schedule"], "active": True, "order": i, "name": t["name"]}
                    if subj:
                        upd["subject"] = subj
                    await teachers.update_one({"id": existing[k]["id"]}, {"$set": upd})
                    result["updated"] += 1
                else:
                    await teachers.insert_one({"id": str(uuid.uuid4()), "active": True, "order": i, **t})
                    result["added"] += 1
            for k, t in existing.items():
                if k not in seen and t.get("active", True):
                    await teachers.update_one({"id": t["id"]}, {"$set": {"active": False}})
                    result["deactivated"] += 1
            result["total"] = len(merged)

        if timing:
            detected = await save_timing_file(timing)
            result["sources"].append("التوقيت")
            if len(detected) >= 8:
                # الصف الأول عادةً الطابور، ثم 4 حصص، فسحة، 3 حصص، فسحة، الحصة الثامنة
                cand = [d for d in detected if _dur(d) == 40][:8]
                if len(cand) == 8:
                    await settings.update_one({"id": "main"}, {"$set": {"period_times": cand}})
                    result["timing_detected"] = 8
        result["settings"] = await get_settings()
        return result

    # ---------------- supervision ----------------
    async def supervision_view():
        doc = await load_supervision()
        tlist = [t async for t in teachers.find({}, {"_id": 0, "id": 1, "name": 1})]
        resolved = resolve_supervision(doc, tlist)
        total = sum(len(d["supervisors"]) + (1 if d["leader"]["name"] else 0) for d in resolved)
        matched = sum(len([s for s in d["supervisors"] if s["teacher_id"]]) + (1 if d["leader"]["teacher_id"] else 0) for d in resolved)
        cfg = await get_settings()
        y0, _ = academic_range(ddate.today().isoformat())
        return {"days": resolved, "updated_at": doc.get("updated_at"), "total": total, "matched": matched,
                "school_name": cfg["school_name"], "year_label": f"{int(y0[:4])}/{int(y0[:4]) + 1}"}

    @router.get("/supervision")
    async def get_supervision(u=Depends(current_user)):
        return await supervision_view()

    @router.put("/supervision")
    async def put_supervision(req: SupervisionReq, u=Depends(current_user)):
        by_day = {d.day: d for d in req.days}
        days_out = [{"day": d, "leader": (by_day[d].leader if d in by_day else "").strip(),
                     "supervisors": [s.strip() for s in (by_day[d].supervisors if d in by_day else []) if s.strip()]} for d in DAYS]
        await supervision.update_one({"id": "main"}, {"$set": {"days": days_out, "updated_at": now_iso()}}, upsert=True)
        return await supervision_view()

    @router.delete("/supervision")
    async def delete_supervision(u=Depends(current_user)):
        await supervision.delete_one({"id": "main"})
        return await supervision_view()

    @router.post("/supervision/import")
    async def import_supervision(file: UploadFile = File(...), u=Depends(current_user)):
        if not (file.filename or "").lower().endswith(".pdf"):
            raise HTTPException(status_code=400, detail="جدول الإشراف: يجب أن يكون ملف PDF")
        tmp = ROOT / "data" / f"_import_{uuid.uuid4().hex}.pdf"
        tmp.write_bytes(await file.read())
        try:
            parsed = parse_supervision_pdf(str(tmp))
        except Exception:
            raise HTTPException(status_code=400, detail="جدول الإشراف: تعذّر قراءة الملف")
        finally:
            tmp.unlink(missing_ok=True)
        if not parsed:
            raise HTTPException(status_code=400, detail="جدول الإشراف: لم يُعثر على أي يوم بقائد إشراف ومشرفين")
        by_day = {p["day"]: p for p in parsed}
        days_out = [{"day": d, "leader": by_day.get(d, {}).get("leader", ""), "supervisors": by_day.get(d, {}).get("supervisors", [])} for d in DAYS]
        await supervision.update_one({"id": "main"}, {"$set": {"days": days_out, "updated_at": now_iso(), "source": file.filename}}, upsert=True)
        out = await supervision_view()
        out["imported_days"] = len(parsed)
        return out

    # ---------------- day ----------------
    @router.get("/days")
    async def list_days(from_date: str, to_date: str, u=Depends(current_user)):
        out = []
        async for d in days.find({"date": {"$gte": from_date, "$lte": to_date}}, {"_id": 0}).sort("date", -1):
            out.append({"date": d["date"], "day_name": d["day_name"], "date_ar": fmt_date_ar(d["date"]), "absent_count": len(d.get("absent", [])),
                        "assignments_count": len([a for a in d.get("assignments", []) if a.get("substitute_id")])})
        return out

    @router.get("/day/{d}")
    async def get_day(d: str, u=Depends(current_user)):
        return await enrich_day(await get_or_new_day(d))

    @router.put("/day/{d}/absent")
    async def set_absent(d: str, req: AbsentReq, u=Depends(current_user)):
        doc = await get_or_new_day(d)
        if not doc["day_name"]:
            raise HTTPException(status_code=400, detail="هذا اليوم ليس يوماً دراسياً")
        ids = list(dict.fromkeys(req.teacher_ids))
        old_ids = doc["absent"]
        was_distributed = any(a.get("substitute_id") for a in doc["assignments"])
        new_ones = [i for i in ids if i not in old_ids]
        late_ids = set(doc.get("late_ids", []))
        if was_distributed:
            late_ids.update(new_ones)
        late_ids &= set(ids)
        doc["absent"] = ids
        doc["late_ids"] = list(late_ids)
        doc["assignments"] = [a for a in doc["assignments"] if a["absent_id"] in ids and a.get("substitute_id") not in ids]
        await save_day(doc)
        return await enrich_day(doc)

    @router.get("/day/{d}/candidates")
    async def candidates(d: str, absent_id: str, period: int, exclude_supervisors: bool = False, u=Depends(current_user)):
        doc = await get_or_new_day(d)
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        return await rank_candidates(doc, tmap, absent_id, period, exclude_busy=False, exclude_supervisors=exclude_supervisors)

    @router.post("/day/{d}/assign")
    async def assign(d: str, req: AssignReq, u=Depends(current_user)):
        doc = await get_or_new_day(d)
        if req.absent_id not in doc["absent"]:
            raise HTTPException(status_code=400, detail="المعلم ليس ضمن الغائبين لهذا اليوم")
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        absent_t = tmap.get(req.absent_id)
        cell = slot_of(absent_t, doc["day_name"], req.period)
        if not cell:
            raise HTTPException(status_code=400, detail="لا توجد حصة للمعلم الغائب في هذا الوقت")
        doc["assignments"] = [a for a in doc["assignments"] if not (a["absent_id"] == req.absent_id and a["period"] == req.period)]
        if req.substitute_id:
            if req.substitute_id in doc["absent"]:
                raise HTTPException(status_code=400, detail="المعلم البديل غائب اليوم")
            if slot_of(tmap.get(req.substitute_id, {}), doc["day_name"], req.period):
                raise HTTPException(status_code=400, detail="المعلم البديل لديه حصة في هذا الوقت")
            if any(a.get("substitute_id") == req.substitute_id and a["period"] == req.period for a in doc["assignments"]):
                raise HTTPException(status_code=400, detail="المعلم البديل مُكلَّف باحتياط آخر في نفس الحصة")
            doc["assignments"].append({"id": str(uuid.uuid4()), "absent_id": req.absent_id, "period": req.period, "class": cell.get("class"),
                                       "subject": cell.get("subject"), "substitute_id": req.substitute_id, "auto": False})
        await save_day(doc)
        return await enrich_day(doc)

    @router.post("/day/{d}/auto")
    async def auto_distribute(d: str, exclude_supervisors: bool = False, u=Depends(current_user)):
        doc = await get_or_new_day(d)
        if not doc["absent"]:
            raise HTTPException(status_code=400, detail="حدّد المعلمين الغائبين أولاً")
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        doc["assignments"] = [a for a in doc["assignments"] if not a.get("auto")]
        slots = []
        for tid in doc["absent"]:
            t = tmap.get(tid)
            if not t:
                continue
            for p in range(1, 9):
                cell = slot_of(t, doc["day_name"], p)
                done = any(a["absent_id"] == tid and a["period"] == p and a.get("substitute_id") for a in doc["assignments"])
                if cell and not done:
                    slots.append((p, tid, cell))
        slots.sort(key=lambda s: s[0])
        filled = 0
        for p, tid, cell in slots:
            cands = await rank_candidates(doc, tmap, tid, p, exclude_busy=True, exclude_supervisors=exclude_supervisors)
            best = pick_fair_candidate(cands)
            if not best:
                continue
            doc["assignments"].append({"id": str(uuid.uuid4()), "absent_id": tid, "period": p, "class": cell.get("class"),
                                       "subject": cell.get("subject"), "substitute_id": best["id"], "auto": True})
            filled += 1
        await save_day(doc)
        res = await enrich_day(doc)
        res["filled"] = filled
        res["unfilled"] = len(slots) - filled
        return res

    @router.delete("/day/{d}")
    async def clear_day(d: str, u=Depends(current_user)):
        await days.delete_one({"date": d})
        return await enrich_day(await get_or_new_day(d))

    @router.delete("/day/{d}/assignments")
    async def clear_assignments(d: str, u=Depends(current_user)):
        doc = await get_or_new_day(d)
        doc["assignments"] = []
        await save_day(doc)
        return await enrich_day(doc)

    # ---------------- stats ----------------
    @router.get("/stats")
    async def stats(from_date: str, to_date: str, teacher_id: Optional[str] = None, u=Depends(current_user)):
        abs_c, sub_c = await load_counts(from_date, to_date)
        p8_c = await load_period_subs(from_date, to_date, 8)
        per_day = []
        day_count = 0
        total_abs = total_sub = total_slots = 0
        async for d in days.find({"date": {"$gte": from_date, "$lte": to_date}}, {"_id": 0}).sort("date", 1):
            if not d.get("absent"):
                continue
            day_count += 1
            n_sub = len([a for a in d.get("assignments", []) if a.get("substitute_id")])
            total_abs += len(d["absent"])
            total_sub += n_sub
            per_day.append({"date": d["date"], "date_ar": fmt_date_ar(d["date"]), "day_name": d["day_name"], "absent": len(d["absent"]), "subs": n_sub})
        per_teacher = []
        async for t in teachers.find({}, {"_id": 0, "schedule": 0}).sort("order", 1):
            per_teacher.append({"id": t["id"], "name": t["name"], "subject": t.get("subject", ""), "quota": t.get("quota", 0), "active": t.get("active", True),
                                "absences": abs_c.get(t["id"], 0), "subs": sub_c.get(t["id"], 0), "subs_p8": p8_c.get(t["id"], 0)})
        by_subject = {}
        for t in per_teacher:
            s = by_subject.setdefault(t["subject"] or "غير محدد", {"subject": t["subject"] or "غير محدد", "absences": 0, "subs": 0, "teachers": 0})
            s["absences"] += t["absences"]; s["subs"] += t["subs"]; s["teachers"] += 1
        top_subs = sorted([t for t in per_teacher if t["subs"]], key=lambda x: -x["subs"])[:5]
        top_abs = sorted([t for t in per_teacher if t["absences"]], key=lambda x: -x["absences"])[:5]
        result = {"from": from_date, "to": to_date, "days": day_count, "total_absences": total_abs, "total_subs": total_sub,
                  "per_day": per_day, "per_teacher": per_teacher, "by_subject": list(by_subject.values()), "top_subs": top_subs, "top_abs": top_abs}
        if teacher_id:
            t = await teachers.find_one({"id": teacher_id}, {"_id": 0, "schedule": 0})
            cfg = await get_settings()
            period_labels = cfg["period_labels"]
            names = {tt["id"]: tt["name"] async for tt in teachers.find({}, {"_id": 0, "id": 1, "name": 1})}
            rows = []
            async for dd in days.find({"date": {"$gte": from_date, "$lte": to_date}}, {"_id": 0}).sort("date", 1):
                for a in dd.get("assignments", []):
                    if a.get("substitute_id") == teacher_id:
                        rows.append({"date": dd["date"], "date_ar": fmt_date_ar(dd["date"]), "day_name": dd["day_name"], "period": a["period"],
                                    "time": period_labels[a["period"] - 1] if a["period"] - 1 < len(period_labels) else "",
                                    "class": a.get("class"), "subject": a.get("subject"), "absent_name": names.get(a["absent_id"], "")})
            result["teacher"] = t
            result["teacher_rows"] = rows
            abs_rows = []
            async for dd in days.find({"date": {"$gte": from_date, "$lte": to_date}, "absent": teacher_id}, {"_id": 0}).sort("date", 1):
                abs_rows.append({"date": dd["date"], "date_ar": fmt_date_ar(dd["date"]), "day_name": dd["day_name"]})
            result["teacher_absence_rows"] = abs_rows
        return result

    @router.delete("/system/reset")
    async def system_reset(u=Depends(current_user)):
        n = await days.count_documents({})
        await days.delete_many({})
        return {"message": "تم تصفير سجلات الغياب والتوزيع بنجاح", "deleted": n}


    @router.get("/day/{d}/export")
    async def export_docx(d: str, absent_id: Optional[str] = None, u=Depends(current_user)):
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        data = await enrich_day(await get_or_new_day(d))
        late_name = None
        if absent_id:
            entry = next((a for a in data["absent"] if a["id"] == absent_id), None)
            late_name = entry["name"] if entry else None
            data = {**data, "assignments": [a for a in data["assignments"] if a["absent_id"] == absent_id]}
        doc = Document()
        sec = doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = Cm(1.2)
        sec.bottom_margin = Cm(1.2)

        def rtl(p):
            pPr = p._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); pPr.append(bidi)

        def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
            p = doc.add_paragraph(); p.alignment = align; rtl(p)
            r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
            r.font.name = "Arial"; r._element.rPr.rFonts.set(qn("w:cs"), "Arial")
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            p.paragraph_format.space_after = Pt(2)
            return p

        logo = ROOT.parent / "frontend" / "public" / "moe-logo.jpeg"
        if logo.exists():
            lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lp.add_run().add_picture(str(logo), width=Cm(3.2))
        para("سلطنة عمان — وزارة التعليم", 11, color="7A1E1E")
        para("المديرية العامة للتعليم بمحافظة شمال الشرقية", 10)
        para(data["school_name"], 14, bold=True)
        if late_name:
            para(f"ملحق إضافي — معلم غائب متأخر: {late_name}", 12, bold=True, color="B45309")
        para(f"توزيع الاحتياط ليوم {data['day_name']} — {data['date_ar']}", 13, bold=True)

        headers = ["م", "المعلم الغائب", "الحصة", "الصف", "المادة", "المعلم البديل", "عدد حصص الاحتياط"]
        table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        bidi = OxmlElement("w:bidiVisual"); tblPr.append(bidi)
        for i, h in enumerate(headers):
            c = table.rows[0].cells[i]; c.text = ""
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(p)
            r = p.add_run(h); r.bold = True; r.font.size = Pt(11)
            shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "E8E8E8"); c._tc.get_or_add_tcPr().append(shd)
        group_start = {}
        for i, a in enumerate(data["assignments"], 1):
            row = table.add_row().cells
            sub_count = str(a["substitute_subs_year"]) if a.get("substitute_id") else ""
            vals = [str(i), a["absent_name"], str(a["period"]), a.get("class") or "", a.get("subject") or "", a["substitute_name"], sub_count]
            for ci, v in enumerate(vals):
                p = row[ci].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(p)
                r = p.add_run(v); r.font.size = Pt(11)
            group_start.setdefault(a["absent_id"], []).append(i)
        for idxs in group_start.values():
            if len(idxs) > 1:
                top, bottom = table.rows[idxs[0]].cells[1], table.rows[idxs[-1]].cells[1]
                merged = top.merge(bottom)
                name = data["assignments"][idxs[0] - 1]["absent_name"]
                for extra in merged.paragraphs[1:]:
                    extra._element.getparent().remove(extra._element)
                merged.paragraphs[0].text = ""
                rr = merged.paragraphs[0].add_run(name); rr.font.size = Pt(11); rr.bold = True
                merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(merged.paragraphs[0])
        if not data["assignments"]:
            row = table.add_row().cells
            row[0].merge(row[-1]); p = row[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(p); p.add_run("لا توجد حصص احتياط لهذا اليوم")
        doc.add_paragraph()
        para("اعتماد إدارة المدرسة", 12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        para("الاسم: .................................        التوقيع: .......................", 11, align=WD_ALIGN_PARAGRAPH.LEFT)
        buf = io.BytesIO(); doc.save(buf); buf.seek(0)
        import urllib.parse
        prefix = f"ملحق {late_name} - " if late_name else ""
        fname_ar = f"{prefix}احتياط {data['day_name']} {data['date_ar'].replace('/', '-')}.docx"
        encoded = urllib.parse.quote(fname_ar)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 headers={"Content-Disposition": f"attachment; filename=substitution_{d}.docx; filename*=UTF-8''{encoded}"})

    # ---------------- swap (تبادل الحصص) ----------------
    async def enrich_swap_day(d: str):
        doc_day = await get_or_new_day(d)
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        cfg = await get_settings()
        period_labels = cfg["period_labels"]
        pairs = []
        async for s in swaps.find({"date": d}, {"_id": 0}).sort("created_at", 1):
            a = tmap.get(s["teacher_a_id"], {})
            b = tmap.get(s["teacher_b_id"], {})
            pairs.append({**s, "teacher_a_name": a.get("name", ""), "teacher_b_name": b.get("name", "")})
        rows = []
        for p in pairs:
            rows.append({"swap_id": p["id"], "period": p["period_b"], "time": period_labels[p["period_b"] - 1], "class": p["class_b"], "subject": p["subject_b"],
                        "original_name": p["teacher_b_name"], "covering_name": p["teacher_a_name"]})
            rows.append({"swap_id": p["id"], "period": p["period_a"], "time": period_labels[p["period_a"] - 1], "class": p["class_a"], "subject": p["subject_a"],
                        "original_name": p["teacher_a_name"], "covering_name": p["teacher_b_name"]})
        return {"date": d, "day_name": doc_day["day_name"], "is_school_day": bool(doc_day["day_name"]), "date_ar": fmt_date_ar(d),
               "school_name": cfg["school_name"], "swaps": pairs, "rows": rows}

    @router.get("/swap/{d}")
    async def get_swap_day(d: str, u=Depends(current_user)):
        return await enrich_swap_day(d)

    @router.get("/swap/{d}/candidates")
    async def swap_candidates(d: str, teacher_id: str, period: int, u=Depends(current_user)):
        doc_day = await get_or_new_day(d)
        if not doc_day["day_name"]:
            raise HTTPException(status_code=400, detail="هذا اليوم ليس يوماً دراسياً")
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        t = tmap.get(teacher_id)
        if not t:
            raise HTTPException(status_code=404, detail="المعلم غير موجود")
        target_class = (slot_of(t, doc_day["day_name"], period) or {}).get("class")
        taken = {a["substitute_id"] for a in doc_day["assignments"] if a["period"] == period and a.get("substitute_id")}
        async for s in swaps.find({"date": d}, {"_id": 0}):
            if s["period_a"] == period:
                taken.add(s["teacher_b_id"])
            if s["period_b"] == period:
                taken.add(s["teacher_a_id"])
        out = []
        for cand in tmap.values():
            if cand["id"] == teacher_id or not cand.get("active", True) or cand["id"] in doc_day["absent"] or cand["id"] in taken:
                continue
            if slot_of(cand, doc_day["day_name"], period):
                continue
            out.append({"id": cand["id"], "name": cand["name"], "subject": cand.get("subject", ""), "same_class": teaches_class(cand, target_class)})
        out.sort(key=lambda c: (not c["same_class"], c["name"]))
        return out

    @router.post("/swap/{d}")
    async def create_swap(d: str, req: SwapCreateReq, u=Depends(current_user)):
        doc_day = await get_or_new_day(d)
        if not doc_day["day_name"]:
            raise HTTPException(status_code=400, detail="هذا اليوم ليس يوماً دراسياً")
        if req.teacher_a_id == req.teacher_b_id:
            raise HTTPException(status_code=400, detail="لا يمكن التبادل مع نفس المعلم")
        tmap = {t["id"]: t async for t in teachers.find({}, {"_id": 0})}
        a, b = tmap.get(req.teacher_a_id), tmap.get(req.teacher_b_id)
        if not a or not b:
            raise HTTPException(status_code=404, detail="أحد المعلمين غير موجود")
        slot_a = slot_of(a, doc_day["day_name"], req.period_a)
        slot_b = slot_of(b, doc_day["day_name"], req.period_b)
        if not slot_a:
            raise HTTPException(status_code=400, detail=f"{a['name']} ليس لديه حصة في الحصة {req.period_a}")
        if not slot_b:
            raise HTTPException(status_code=400, detail=f"{b['name']} ليس لديه حصة في الحصة {req.period_b}")
        if req.period_a != req.period_b:
            if slot_of(b, doc_day["day_name"], req.period_a):
                raise HTTPException(status_code=400, detail=f"{b['name']} مشغول في الحصة {req.period_a}")
            if slot_of(a, doc_day["day_name"], req.period_b):
                raise HTTPException(status_code=400, detail=f"{a['name']} مشغول في الحصة {req.period_b}")
        swap_doc = {"id": str(uuid.uuid4()), "date": d, "day_name": doc_day["day_name"],
               "teacher_a_id": a["id"], "period_a": req.period_a, "class_a": slot_a.get("class"), "subject_a": slot_a.get("subject"),
               "teacher_b_id": b["id"], "period_b": req.period_b, "class_b": slot_b.get("class"), "subject_b": slot_b.get("subject"),
               "created_at": now_iso()}
        await swaps.insert_one(swap_doc)
        return await enrich_swap_day(d)

    @router.delete("/swap/{d}/{swap_id}")
    async def delete_swap(d: str, swap_id: str, u=Depends(current_user)):
        await swaps.delete_one({"id": swap_id, "date": d})
        return await enrich_swap_day(d)

    @router.delete("/swap/{d}")
    async def clear_swap_day(d: str, u=Depends(current_user)):
        await swaps.delete_many({"date": d})
        return await enrich_swap_day(d)

    @router.get("/swap/{d}/export")
    async def export_swap_docx(d: str, u=Depends(current_user)):
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        data = await enrich_swap_day(d)
        doc = Document()
        sec = doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = Cm(1.2)
        sec.bottom_margin = Cm(1.2)

        def rtl(p):
            pPr = p._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); pPr.append(bidi)

        def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
            p = doc.add_paragraph(); p.alignment = align; rtl(p)
            r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
            r.font.name = "Arial"; r._element.rPr.rFonts.set(qn("w:cs"), "Arial")
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            p.paragraph_format.space_after = Pt(2)
            return p

        logo = ROOT.parent / "frontend" / "public" / "moe-logo.jpeg"
        if logo.exists():
            lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lp.add_run().add_picture(str(logo), width=Cm(3.2))
        para("سلطنة عمان — وزارة التعليم", 11, color="7A1E1E")
        para("المديرية العامة للتعليم بمحافظة شمال الشرقية", 10)
        para(data["school_name"], 14, bold=True)
        para(f"تبادل الحصص ليوم {data['day_name']} — {data['date_ar']}", 13, bold=True)

        headers = ["م", "المعلم", "الحصة", "الصف", "المادة", "يُغطّيها"]
        table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        bidi = OxmlElement("w:bidiVisual"); tblPr.append(bidi)
        for i, h in enumerate(headers):
            c = table.rows[0].cells[i]; c.text = ""
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(p)
            r = p.add_run(h); r.bold = True; r.font.size = Pt(11)
            shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "E8E8E8"); c._tc.get_or_add_tcPr().append(shd)
        rows = data["rows"]
        for i, rr in enumerate(rows, 1):
            row = table.add_row().cells
            vals = [str(i), rr["original_name"], str(rr["period"]), rr.get("class") or "", rr.get("subject") or "", rr["covering_name"]]
            for ci, v in enumerate(vals):
                p = row[ci].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(p)
                r = p.add_run(v); r.font.size = Pt(11)
        if not rows:
            row = table.add_row().cells
            row[0].merge(row[-1]); p = row[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rtl(p); p.add_run("لا توجد عمليات تبادل لهذا اليوم")
        doc.add_paragraph()
        para("اعتماد إدارة المدرسة", 12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        para("الاسم: .................................        التوقيع: .......................", 11, align=WD_ALIGN_PARAGRAPH.LEFT)
        buf = io.BytesIO(); doc.save(buf); buf.seek(0)
        import urllib.parse
        fname_ar = f"تبادل حصص {data['day_name']} {data['date_ar'].replace('/', '-')}.docx"
        encoded = urllib.parse.quote(fname_ar)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 headers={"Content-Disposition": f"attachment; filename=swap_{d}.docx; filename*=UTF-8''{encoded}"})

    return router
